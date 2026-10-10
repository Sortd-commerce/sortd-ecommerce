"""Second-pass web research: prices, images, ingredients (Carrefour + Open Food Facts)."""

from __future__ import annotations

import html as html_lib
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "Sortd_Launch_Range_internet_enriched.xlsx"
SHEET = "Launch range"
DATA_ROW_START = 5
KEY_ROW = 3
TIMEOUT_MS = 55000


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", html_lib.unescape(s or "").lower()).strip()


def _score(query: str, candidate: str) -> int:
    q = set(_norm(query).split())
    c = set(_norm(candidate).split())
    return len(q & c) if q else 0


def _build_query(row: dict) -> str:
    brand = (row.get("brand") or "").strip()
    name = (row.get("product_name") or "").strip()
    variant = (row.get("variant") or "").strip()
    pack = (row.get("pack_line") or "").strip()
    if variant and variant.lower() not in name.lower():
        name = f"{name} {variant}".strip()
    size = ""
    if pack:
        m = re.search(r"(\d+\s*(?:g|kg|ml|l|pcs))", pack, re.I)
        size = m.group(1) if m else ""
    return " ".join(p for p in (brand, name, size) if p).strip()


def _off_lookup(query: str) -> dict:
    url = (
        "https://world.openfoodfacts.org/cgi/search.pl?"
        + urllib.parse.urlencode(
            {"search_terms": query, "search_simple": 1, "action": "process", "json": 1, "page_size": 5}
        )
    )
    try:
        raw = urllib.request.urlopen(
            urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}),
            timeout=12,
        ).read()
        data = json.loads(raw)
    except (urllib.error.URLError, json.JSONDecodeError, TimeoutError):
        return {}
    for hit in data.get("products") or []:
        if _score(query, hit.get("product_name") or "") < 2:
            continue
        out: dict = {}
        if hit.get("ingredients_text"):
            out["ingredients_full"] = hit["ingredients_text"]
        if hit.get("image_url"):
            out["main_image"] = hit["image_url"]
        return out
    return {}


def _parse_listings(text: str) -> list[tuple[str, float]]:
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    listings: list[tuple[str, float]] = []
    i = 0
    while i < len(lines):
        if lines[i] == "AED" and i + 1 < len(lines):
            try:
                price = float(lines[i + 1].replace(",", ""))
            except ValueError:
                i += 1
                continue
            title = ""
            for back in range(1, 7):
                if i - back < 0:
                    break
                cand = lines[i - back]
                if cand in {"AED", "BESTSELLER", "Bright Bites", "Choose Better"}:
                    continue
                if cand.startswith("Tod.") or "per Kilo" in cand or re.match(r"^\d+% OFF$", cand):
                    continue
                if re.match(r"^[\d.]+$", cand):
                    continue
                if len(cand) > 10:
                    title = cand
                    break
            if title:
                listings.append((title, price))
            i += 2
            continue
        i += 1
    return listings


def _best_listing(query: str, listings: list[tuple[str, float]]) -> tuple[str, float] | None:
    best: tuple[str, float] | None = None
    best_score = 0
    for title, price in listings:
        sc = _score(query, title)
        if sc > best_score:
            best_score = sc
            best = (title, price)
    if best_score >= 3:
        return best
    if best_score >= 2:
        return best
    return None


def _is_empty(v) -> bool:
    return v in (None, "")


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    keys = [c.value for c in ws[KEY_ROW]]
    key_to_col = {k: i + 1 for i, k in enumerate(keys) if k}

    targets: list[tuple[int, dict, str]] = []
    for r in range(DATA_ROW_START, ws.max_row + 1):
        sku = ws.cell(r, key_to_col["sortd_sku"]).value
        if not sku:
            continue
        row = {k: ws.cell(r, key_to_col[k]).value for k in key_to_col}
        needs = (
            _is_empty(row.get("price"))
            or _is_empty(row.get("main_image"))
            or _is_empty(row.get("ingredients_full"))
            or (row.get("price") and float(row.get("price")) > 120 and (row.get("aisle") or "").startswith("Snacks"))
        )
        if needs:
            targets.append((r, row, str(sku)))

    print(f"Round 2 targets: {len(targets)}", flush=True)
    updated = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )
        page = browser.new_page(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )
        page.set_default_timeout(TIMEOUT_MS)

        for idx, (r, row, sku) in enumerate(targets, start=1):
            query = _build_query(row)
            patch: dict = {}

            off = _off_lookup(query)
            if _is_empty(row.get("ingredients_full")) and off.get("ingredients_full"):
                patch["ingredients_full"] = off["ingredients_full"]
            if _is_empty(row.get("main_image")) and off.get("main_image"):
                patch["main_image"] = off["main_image"]

            url = "https://www.carrefouruae.com/mafuae/en/search?keyword=" + urllib.parse.quote(query)
            search_ok = False
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=25000)
                page.wait_for_timeout(2200)
                search_ok = True
                text = page.evaluate("() => document.body.innerText")
                listings = _parse_listings(text)
                best = _best_listing(query, listings)
                if best:
                    title, price = best
                    if _is_empty(row.get("price")) or (
                        row.get("price") and float(row.get("price")) > 120 and price < float(row.get("price")) * 0.5
                    ):
                        patch["price"] = round(price, 2)
                    # Product link by partial title match
                    href = page.evaluate(
                        """(title) => {
                          const links = [...document.querySelectorAll('a[href*="/p/"]')];
                          const t = title.toLowerCase();
                          for (const a of links) {
                            const name = (a.textContent || '').trim().toLowerCase();
                            if (name && (name.includes(t.slice(0,20)) || t.includes(name.slice(0,20)))) return a.href;
                          }
                          return links[0]?.href || '';
                        }""",
                        title,
                    )
                    if href and search_ok and _is_empty(row.get("main_image")):
                        page.goto(href, wait_until="domcontentloaded", timeout=20000)
                        page.wait_for_timeout(1500)
                        meta = page.evaluate(
                            """() => ({
                              image: document.querySelector('meta[property=\"og:image\"]')?.content || '',
                              desc: document.querySelector('meta[property=\"og:description\"]')?.content || ''
                            })"""
                        )
                        if meta.get("image"):
                            patch["main_image"] = meta["image"]
                        if _is_empty(row.get("short_desc")) and meta.get("desc"):
                            patch["short_desc"] = meta["desc"][:900]
            except Exception as exc:  # noqa: BLE001
                print(f"{sku}: carrefour error {exc}", flush=True)

            applied = 0
            for k, v in patch.items():
                if k not in key_to_col or _is_empty(v):
                    continue
                cur = ws.cell(r, key_to_col[k]).value
                if _is_empty(cur) or (k == "price" and v != cur):
                    ws.cell(r, key_to_col[k]).value = v
                    applied += 1
            if applied:
                updated += 1
                wb.save(WORKBOOK)
                if "nutrition_note" in key_to_col:
                    note = ws.cell(r, key_to_col["nutrition_note"]).value or ""
                    suffix = "Round-2 web enrichment."
                    if suffix not in str(note):
                        ws.cell(r, key_to_col["nutrition_note"]).value = (
                            (str(note) + " " if note else "") + suffix
                        ).strip()
                print(f"[{idx}/{len(targets)}] {sku}: +{applied} ({', '.join(patch.keys())})", flush=True)

        browser.close()

    wb.save(WORKBOOK)
    print(f"Round 2 done. Rows updated: {updated}. Saved {WORKBOOK}", flush=True)


if __name__ == "__main__":
    main()
