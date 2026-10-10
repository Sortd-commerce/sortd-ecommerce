"""Fill empty main_image from Carrefour UAE search result thumbnails."""

from __future__ import annotations

import re
import urllib.parse
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "Sortd_Launch_Range_internet_enriched.xlsx"
SHEET = "Launch range"
DATA_ROW_START = 5
KEY_ROW = 3


def _build_query(row: dict) -> str:
    brand = (row.get("brand") or "").strip()
    name = (row.get("product_name") or "").strip()
    variant = (row.get("variant") or "").strip()
    if variant and variant.lower() not in name.lower():
        name = f"{name} {variant}".strip()
    pack = (row.get("pack_line") or "").strip()
    size = ""
    if pack:
        m = re.search(r"(\d+\s*(?:g|kg|ml|l|pcs))", pack, re.I)
        size = m.group(1) if m else ""
    return " ".join(p for p in (brand, name, size) if p).strip()


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    keys = [c.value for c in ws[KEY_ROW]]
    key_to_col = {k: i + 1 for i, k in enumerate(keys) if k}
    targets = []
    for r in range(DATA_ROW_START, ws.max_row + 1):
        sku = ws.cell(r, key_to_col["sortd_sku"]).value
        if not sku:
            continue
        if ws.cell(r, key_to_col["main_image"]).value:
            continue
        row = {k: ws.cell(r, key_to_col[k]).value for k in key_to_col}
        targets.append((r, row, str(sku)))

    filled = 0
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--disable-blink-features=AutomationControlled"])
        page = browser.new_page(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )
        for r, row, sku in targets:
            query = _build_query(row)
            url = "https://www.carrefouruae.com/mafuae/en/search?keyword=" + urllib.parse.quote(query)
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=25000)
                page.wait_for_timeout(2000)
                img = page.evaluate(
                    """() => {
                      const imgs = [...document.querySelectorAll('img[src]')];
                      for (const el of imgs) {
                        const src = el.src || '';
                        if (src.includes('product') || src.includes('cdn') && src.includes('carrefour'))
                          return src;
                      }
                      return imgs.find(i => i.width >= 80)?.src || '';
                    }"""
                )
                if img and img.startswith("http"):
                    ws.cell(r, key_to_col["main_image"]).value = img
                    wb.save(WORKBOOK)
                    filled += 1
                    print(sku, img[:80], flush=True)
            except Exception as exc:  # noqa: BLE001
                print(sku, "err", exc, flush=True)
        browser.close()
    print(f"Filled {filled} images")


if __name__ == "__main__":
    main()
