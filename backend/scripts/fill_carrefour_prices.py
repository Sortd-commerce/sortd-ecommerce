"""Fill empty price cells using Carrefour UAE search (Playwright)."""

from __future__ import annotations

import html as html_lib
import re
import urllib.parse
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "Sortd_Launch_Range_internet_enriched.xlsx"
SHEET = "Launch range"
DATA_ROW_START = 5
KEY_ROW = 3


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", html_lib.unescape(s or "").lower()).strip()


def _score(query: str, title: str) -> int:
    q = set(_norm(query).split())
    t = set(_norm(title).split())
    return len(q & t) if q else 0


def _parse_listings(text: str) -> list[tuple[str, float]]:
    """Parse Carrefour search page innerText into (title, price AED) pairs."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    listings: list[tuple[str, float]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line == "AED" and i + 1 < len(lines):
            try:
                price = float(lines[i + 1].replace(",", ""))
            except ValueError:
                i += 1
                continue
            # title is usually a few lines above, skip schedule/noise
            title = ""
            for back in range(1, 6):
                if i - back < 0:
                    break
                cand = lines[i - back]
                if cand in {"AED", "Tod. 2:30 PM", "BESTSELLER", "Bright Bites", "Choose Better"}:
                    continue
                if re.match(r"^\d+% OFF$", cand):
                    continue
                if "per Kilo" in cand or "per Liter" in cand:
                    continue
                if re.match(r"^[\d.]+$", cand):
                    continue
                if len(cand) > 8:
                    title = cand
                    break
            if title:
                listings.append((title, price))
            i += 2
            continue
        i += 1
    return listings


def _best_price(query: str, listings: list[tuple[str, float]]) -> float | None:
    best_score = 0
    best_price: float | None = None
    brand_hint = query.split()[0:2]
    for title, price in listings:
        score = _score(query, title)
        if score > best_score:
            best_score = score
            best_price = price
    if best_score >= 3:
        return best_price
    # fallback: title contains brand + at least one product token
    q_norm = _norm(query)
    for title, price in listings:
        t_norm = _norm(title)
        if all(w in t_norm for w in _norm(" ".join(brand_hint)).split() if len(w) > 2):
            if _score(query, title) >= 2:
                return price
    if best_score >= 2:
        return best_price
    return None


def _build_query(row: dict) -> str:
    brand = (row.get("brand") or "").strip()
    name = (row.get("product_name") or "").strip()
    variant = (row.get("variant") or "").strip()
    pack = (row.get("pack_line") or "").strip()
    if variant and variant.lower() not in name.lower():
        name = f"{name} {variant}".strip()
    # Carrefour search works best with brand + title + size token, not full pack line.
    size = ""
    if pack:
        m = re.search(r"(\d+\s*(?:g|kg|ml|l|pcs))", pack, re.I)
        size = m.group(1) if m else pack
    return " ".join(p for p in (brand, name, size) if p).strip()


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    keys = [c.value for c in ws[KEY_ROW]]
    key_to_col = {k: i + 1 for i, k in enumerate(keys) if k}

    rows: list[tuple[int, dict, str]] = []
    for r in range(DATA_ROW_START, ws.max_row + 1):
        sku = ws.cell(r, key_to_col["sortd_sku"]).value
        if not sku:
            continue
        row = {k: ws.cell(r, key_to_col[k]).value for k in key_to_col}
        if row.get("price") not in (None, ""):
            continue
        rows.append((r, row, str(sku)))

    filled = 0
    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )
        page = browser.new_page(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )
        page.set_default_timeout(60000)
        for r, row, sku in rows:
            query = _build_query(row)
            url = "https://www.carrefouruae.com/mafuae/en/search?keyword=" + urllib.parse.quote(query)
            try:
                for attempt in range(2):
                    try:
                        page.goto(url, wait_until="domcontentloaded")
                        break
                    except Exception:
                        if attempt == 0:
                            page.wait_for_timeout(1500)
                            continue
                        raise
                page.wait_for_timeout(3000)
                text = page.evaluate("() => document.body.innerText")
                listings = _parse_listings(text)
                price = _best_price(query, listings)
                if price:
                    ws.cell(r, key_to_col["price"]).value = round(price, 2)
                    note = row.get("nutrition_note") or ""
                    suffix = "Price from Carrefour UAE search (verify before Live)."
                    if suffix not in str(note):
                        ws.cell(r, key_to_col["nutrition_note"]).value = (
                            (str(note) + " " if note else "") + suffix
                        ).strip()
                    status = row.get("status")
                    if status in ("Needs price", "Paused", "Draft"):
                        ws.cell(r, key_to_col["status"]).value = "Ready for review"
                    filled += 1
                    print(f"{sku}: AED {price}", flush=True)
                else:
                    print(f"{sku}: no Carrefour match", flush=True)
            except Exception as exc:  # noqa: BLE001
                print(f"{sku}: error {exc}", flush=True)
        browser.close()

    wb.save(WORKBOOK)
    print(f"Filled {filled} prices. Saved {WORKBOOK}")


if __name__ == "__main__":
    main()
