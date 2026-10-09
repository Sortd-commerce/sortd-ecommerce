"""Import products from the Sortd product listing Excel workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from catalog.image_fetch import ImageFetchError, is_media_ref, is_remote_ref, load_image, load_pdf, normalize_image_url, split_image_urls
from catalog.reports import publish_report
from catalog.images import sync_image_order
from catalog.models import (
    Allergen,
    Category,
    Ingredient,
    LabReport,
    NutritionFact,
    NutritionProfile,
    Product,
    ProductImage,
    ProductStatus,
    ProductVariant,
    RelatedKind,
    RelatedProduct,
)
from catalog.writer import CatalogWriteError

SKU_RE = re.compile(r"^SRT-[A-Z0-9-]+$", re.IGNORECASE)
SKIP_IMPORT_SKUS = frozenset({"SRT-EXAMPLE-001"})
INSTRUCTION_MARKERS = {
    "sortd fills",
    "from supplier",
    "unique code",
    "shown as the big title",
    "pick from the list",
    "top-level shop aisle",
    "reference only",
    "auto",
}

AISLE_TO_CATEGORY = {
    "Breakfast & spreads": ("Breakfast & spreads", "breakfast-spreads", 1),
    "Snacks & bars": ("Snacks & bars", "snacks-bars", 2),
    "Chocolate": ("Chocolate", "chocolate", 3),
    "Drinks": ("Drinks & hydration", "drinks-hydration", 4),
    "Pantry": ("Cooking & pantry", "cooking-pantry", 5),
}

EXCEL_STATUS_TO_DB = {
    "live": ProductStatus.ACTIVE,
    "paused": ProductStatus.ARCHIVED,
}

LABEL_TEMPLATE_HIGHLIGHTS: dict[str, list[str]] = {
    "Protein & snack bars": ["protein_g", "carbs_g", "fat_g"],
    "Savoury snacks": ["protein_g", "fat_g", "sodium_mg"],
    "Cooking oils": ["fat_g", "trans_fat_g", "vitamin_e_mg"],
    "Drinks & juices": ["sugars_g", "potassium_mg", "sodium_mg"],
    "Chocolate": ["cacao_pct", "sugars_g", "fat_g"],
    "Pulses, flour & pasta": ["protein_g", "carbs_g", "fibre_g"],
    "Nut butters & spreads": ["protein_g", "fat_g", "sugars_g"],
    "Breakfast & oats": ["protein_g", "fibre_g", "sugars_g"],
    "Sauces & condiments": ["sugars_g", "sodium_mg", "fat_g"],
}

NUTRITION_FIELDS: dict[str, tuple[str, str, str | None]] = {
    "energy_kcal": ("Energy", "kcal", None),
    "protein_g": ("Protein", "g", None),
    "carbs_g": ("Carbohydrate", "g", None),
    "sugars_g": ("Total sugars", "g", "lv_sugars_level"),
    "added_sugars_g": ("Added sugars", "g", None),
    "fibre_g": ("Fibre", "g", None),
    "fat_g": ("Fat", "g", "lv_fat_level"),
    "saturates_g": ("Saturates", "g", "lv_saturates_level"),
    "mufa_g": ("MUFA", "g", None),
    "pufa_g": ("PUFA", "g", None),
    "trans_fat_g": ("Trans fat", "g", None),
    "cholesterol_mg": ("Cholesterol", "mg", None),
    "sodium_mg": ("Sodium", "mg", None),
    "potassium_mg": ("Potassium", "mg", None),
    "vitamin_e_mg": ("Vitamin E", "mg", None),
    "cacao_pct": ("Cacao / cocoa solids", "%", None),
}

LEVEL_MAP = {
    "low": "low",
    "medium": "medium",
    "high": "high",
}


@dataclass
class ImportIssue:
    row: int
    sku: str
    field: str
    message: str
    level: str = "error"


@dataclass
class ExcelImportResult:
    valid: bool
    dry_run: bool
    row_count: int
    created: int = 0
    updated: int = 0
    skipped: int = 0
    aisle_images_updated: int = 0
    issues: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)

    @property
    def errors(self) -> list[ImportIssue]:
        return [issue for issue in self.issues if issue.level == "error"]


class ExcelImportError(ValueError):
    pass


class ExcelCatalogImporter:
    KEY_ROW = 3
    DATA_START_ROW = 4
    SHEET_NAME = "Launch range"
    AISLES_SHEET = "Aisles"
    BATCH_SIZE = 200
    _media_root: Path | None = None

    def import_file(
        self,
        file_obj,
        *,
        dry_run: bool = False,
        force_active: bool = False,
        skip_images: bool = False,
        media_root: Path | None = None,
    ) -> ExcelImportResult:
        self._skip_images = skip_images
        self._media_root = media_root or self._workbook_media_root(file_obj)
        workbook = self._load_workbook(file_obj)
        rows = self._parse_product_rows(workbook)
        aisle_rows = self._parse_aisle_rows(workbook)
        result = self.validate_rows(rows)
        self._validate_aisle_rows(aisle_rows, result)
        result.dry_run = dry_run
        if dry_run or result.errors:
            return result

        try:
            with transaction.atomic():
                category_map = self._ensure_categories()
                if not skip_images:
                    self._import_aisle_images(category_map, aisle_rows, result)
                parent_groups = self._bulk_import_rows(rows, category_map, result, force_active=force_active)
                self._bulk_link_parent_groups(parent_groups)
        except (CatalogWriteError, ImageFetchError, ExcelImportError) as exc:
            raise ExcelImportError(str(exc)) from exc

        return result

    def parse_workbook(self, file_obj, *, media_root: Path | None = None) -> list[dict[str, Any]]:
        self._media_root = media_root or self._workbook_media_root(file_obj)
        return self._parse_product_rows(self._load_workbook(file_obj))

    @staticmethod
    def _workbook_media_root(file_obj) -> Path | None:
        name = getattr(file_obj, "name", None)
        if not name or str(name).startswith("<"):
            return None
        path = Path(str(name))
        if path.suffix.lower() == ".xlsx" and path.parent.exists():
            return path.parent.resolve()
        return None

    def _load_workbook(self, file_obj):
        if hasattr(file_obj, "read"):
            payload = file_obj.read()
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
            return load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
        return load_workbook(filename=file_obj, read_only=True, data_only=True)

    def _parse_product_rows(self, workbook) -> list[dict[str, Any]]:
        if self.SHEET_NAME not in workbook.sheetnames:
            raise ExcelImportError(f"Worksheet '{self.SHEET_NAME}' was not found in the workbook.")
        sheet = workbook[self.SHEET_NAME]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < self.DATA_START_ROW:
            raise ExcelImportError("Worksheet does not contain any product rows.")

        keys = [self._cell_text(value) for value in rows[self.KEY_ROW - 1]]
        parsed: list[dict[str, Any]] = []
        for index, values in enumerate(rows[self.DATA_START_ROW - 1 :], start=self.DATA_START_ROW):
            if not values or all(value in (None, "") for value in values):
                continue
            row = {"__row": index}
            for key_index, key in enumerate(keys):
                if not key:
                    continue
                row[key] = values[key_index] if key_index < len(values) else None
            if self._is_product_row(row):
                parsed.append(row)
        return parsed

    def _parse_aisle_rows(self, workbook) -> list[dict[str, Any]]:
        if self.AISLES_SHEET not in workbook.sheetnames:
            return []
        sheet = workbook[self.AISLES_SHEET]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        headers = [self._cell_text(value).lower() for value in rows[0]]
        try:
            aisle_index = headers.index("aisle")
            image_index = headers.index("image_url")
        except ValueError:
            return []

        parsed: list[dict[str, Any]] = []
        for index, values in enumerate(rows[1:], start=2):
            if not values or all(value in (None, "") for value in values):
                continue
            aisle = self._cell_text(values[aisle_index] if aisle_index < len(values) else "")
            image_url = self._cell_text(values[image_index] if image_index < len(values) else "")
            if not aisle and not image_url:
                continue
            parsed.append({"__row": index, "aisle": aisle, "image_url": image_url})
        return parsed

    def _validate_aisle_rows(self, aisle_rows: list[dict[str, Any]], result: ExcelImportResult) -> None:
        for row in aisle_rows:
            row_no = int(row["__row"])
            aisle = self._cell_text(row.get("aisle"))
            image_url = self._cell_text(row.get("image_url"))
            if not aisle:
                self._add_issue(result, row_no, "", "aisle", "Aisle name is required on the Aisles sheet.")
                continue
            if aisle not in AISLE_TO_CATEGORY:
                self._add_issue(result, row_no, aisle, "aisle", f"Unknown aisle '{aisle}'.")
                continue
            if image_url and not self._is_media_ref(image_url):
                self._add_warning(
                    result,
                    row_no,
                    aisle,
                    "image_url",
                    "Use a public https URL or a path relative to the workbook folder.",
                )

    def _import_aisle_images(
        self,
        category_map: dict[str, Category],
        aisle_rows: list[dict[str, Any]],
        result: ExcelImportResult,
    ) -> None:
        for row in aisle_rows:
            aisle = self._cell_text(row.get("aisle"))
            image_url = self._cell_text(row.get("image_url"))
            if not aisle or aisle not in AISLE_TO_CATEGORY or not image_url:
                continue
            if not self._is_media_ref(image_url):
                continue
            slug = AISLE_TO_CATEGORY[aisle][1]
            category = category_map[slug]
            if self._set_category_image(category, image_url):
                result.aisle_images_updated += 1

    def validate_rows(self, rows: list[dict[str, Any]]) -> ExcelImportResult:
        result = ExcelImportResult(valid=True, dry_run=False, row_count=len(rows))
        seen_skus: dict[str, int] = {}

        for row in rows:
            row_no = int(row["__row"])
            sku = self._cell_text(row.get("sortd_sku"))
            if not sku:
                self._add_issue(result, row_no, "", "sortd_sku", "Sortd SKU is required.")
                continue
            if sku in seen_skus:
                self._add_issue(
                    result,
                    row_no,
                    sku,
                    "sortd_sku",
                    f"Duplicate Sortd SKU in file (first seen on row {seen_skus[sku]}).",
                )
            else:
                seen_skus[sku] = row_no

            if not SKU_RE.match(sku):
                self._add_issue(result, row_no, sku, "sortd_sku", "Sortd SKU must look like SRT-SNK-021.")

            product_name = self._cell_text(row.get("product_name"))
            if not product_name:
                self._add_issue(result, row_no, sku, "product_name", "Product name is required.")

            if not self._cell_text(row.get("brand")):
                self._add_issue(result, row_no, sku, "brand", "Brand is required.")

            aisle = self._cell_text(row.get("aisle"))
            if not aisle:
                self._add_issue(result, row_no, sku, "aisle", "Aisle is required.")
            elif aisle not in AISLE_TO_CATEGORY:
                self._add_issue(result, row_no, sku, "aisle", f"Unknown aisle '{aisle}'.")

            status = self._excel_status(row)
            if status is None:
                self._add_issue(result, row_no, sku, "status", f"Unsupported status '{self._cell_text(row.get('status'))}'.")

            price = row.get("price")
            if status == ProductStatus.ACTIVE and not self._has_price(price):
                self._add_issue(result, row_no, sku, "price", "Live products need a selling price.")

            if status == ProductStatus.ACTIVE and not self._cell_text(row.get("short_desc")):
                self._add_issue(result, row_no, sku, "short_desc", "Live products need a short description.")

            if status == ProductStatus.ACTIVE and not self._variant_title(row):
                self._add_issue(
                    result,
                    row_no,
                    sku,
                    "pack_line",
                    "Live products need a pack line or net quantity + unit.",
                )

            main_image = self._cell_text(row.get("main_image"))
            if status == ProductStatus.ACTIVE and not main_image:
                self._add_issue(result, row_no, sku, "main_image", "Live products need a main image URL.")
            elif main_image:
                self._validate_image_reference(
                    result,
                    row_no,
                    sku,
                    "main_image",
                    main_image,
                    required=status == ProductStatus.ACTIVE,
                )

            for index, url in enumerate(split_image_urls(self._cell_text(row.get("gallery"))), start=1):
                self._validate_image_reference(
                    result,
                    row_no,
                    sku,
                    "gallery",
                    url,
                    required=False,
                    label=f"Gallery image {index}",
                )

            lab_report_url = self._cell_text(row.get("lab_report_url"))
            if lab_report_url:
                self._validate_image_reference(
                    result,
                    row_no,
                    sku,
                    "lab_report_url",
                    lab_report_url,
                    required=False,
                    label="Lab report PDF",
                )

            if self._has_price(price):
                try:
                    Decimal(str(price))
                except (InvalidOperation, TypeError):
                    self._add_issue(result, row_no, sku, "price", "Selling price must be a number.")

            compare_price = row.get("compare_price")
            if compare_price not in (None, "") and not self._has_price(compare_price):
                self._add_warning(result, row_no, sku, "compare_price", "Compare-at price is not a number and will be ignored.")

            ingredients = self._ingredients_from_row(row)
            if status == ProductStatus.ACTIVE and not ingredients and not self._cell_text(row.get("ingredients_full")):
                self._add_issue(result, row_no, sku, "ingredients", "Live products need ingredients.")

        result.valid = not result.errors
        return result

    def _bulk_import_rows(
        self,
        rows: list[dict[str, Any]],
        categories: dict[str, Category],
        result: ExcelImportResult,
        *,
        force_active: bool = False,
    ) -> dict[str, list[str]]:
        skus = [self._cell_text(row.get("sortd_sku")) for row in rows]
        slugs = [self._product_slug(row) for row in rows]
        variant_map = {
            variant.sku.upper(): variant
            for variant in ProductVariant.objects.select_related("product").filter(sku__in=skus)
        }
        product_by_slug = {product.slug: product for product in Product.objects.filter(slug__in=slugs)}

        new_products: list[Product] = []
        update_products: list[Product] = []
        row_meta: list[dict[str, Any]] = []
        parent_groups: dict[str, list[str]] = {}

        for row in rows:
            sku = self._cell_text(row.get("sortd_sku"))
            slug = self._product_slug(row)
            aisle = self._cell_text(row.get("aisle"))
            category = categories[AISLE_TO_CATEGORY[aisle][1]]
            title = self._product_title(row)
            description = self._cell_text(row.get("short_desc"))
            brand = self._cell_text(row.get("brand"))
            shelf = self._cell_text(row.get("shelf"))
            tags = self._cell_text(row.get("tags"))
            status = ProductStatus.ACTIVE if force_active else (self._excel_status(row) or ProductStatus.DRAFT)

            variant = variant_map.get(sku.upper())
            product = variant.product if variant else product_by_slug.get(slug)
            created = product is None
            if product is None:
                product = Product(
                    title=title,
                    slug=slug,
                    brand=brand,
                    description=description,
                    category=category,
                    shelf=shelf,
                    tags=tags,
                    status=status,
                )
                new_products.append(product)
                product_by_slug[slug] = product
            else:
                product.title = title
                product.brand = brand
                product.description = description
                product.category = category
                product.shelf = shelf
                product.tags = tags
                product.status = status
                update_products.append(product)

            row_meta.append({"row": row, "sku": sku, "slug": slug, "product": product, "created": created})
            group = self._cell_text(row.get("parent_group"))
            if group:
                parent_groups.setdefault(group, []).append(slug)

        if new_products:
            Product.objects.bulk_create(new_products, batch_size=self.BATCH_SIZE)
        if update_products:
            Product.objects.bulk_update(
                update_products,
                ["title", "brand", "description", "category", "shelf", "tags", "status", "updated_at"],
                batch_size=self.BATCH_SIZE,
            )

        new_variants: list[ProductVariant] = []
        update_variants: list[ProductVariant] = []
        for meta in row_meta:
            row = meta["row"]
            product = meta["product"]
            sku = meta["sku"]
            on_hand = max(int(row.get("stock") or 0), 0)
            fields = {
                "product": product,
                "title": self._variant_title(row),
                "price": self._decimal_or_zero(row.get("price")),
                "compare_at_price": self._optional_decimal(row.get("compare_price")),
                "unit_count": max(int(row.get("units_in_pack") or 1), 1),
                "max_order": self._optional_positive_int(row.get("max_order")),
                "on_hand": on_hand,
                "is_active": True,
            }
            variant = variant_map.get(sku.upper())
            if variant is None:
                new_variants.append(ProductVariant(sku=sku, **fields))
                continue
            if variant.product_id and variant.product_id != product.id:
                raise ExcelImportError(f"SKU '{sku}' already belongs to another product.")
            for key, value in fields.items():
                setattr(variant, key, value)
            update_variants.append(variant)

        if new_variants:
            ProductVariant.objects.bulk_create(new_variants, batch_size=self.BATCH_SIZE)
        if update_variants:
            ProductVariant.objects.bulk_update(
                update_variants,
                ["product", "title", "price", "compare_at_price", "unit_count", "max_order", "on_hand", "is_active"],
                batch_size=self.BATCH_SIZE,
            )

        self._bulk_replace_labels(row_meta)

        for meta in row_meta:
            if not getattr(self, "_skip_images", False):
                self._import_images(meta["product"], meta["row"])
            self._import_lab_report(meta["product"], meta["row"])
            if meta["created"]:
                result.created += 1
            else:
                result.updated += 1

        return parent_groups

    def _bulk_replace_labels(self, row_meta: list[dict[str, Any]]) -> None:
        product_ids = [meta["product"].id for meta in row_meta if meta["product"].id]
        if not product_ids:
            return

        NutritionProfile.objects.filter(product_id__in=product_ids).delete()
        Ingredient.objects.filter(product_id__in=product_ids).delete()
        Allergen.objects.filter(product_id__in=product_ids).delete()

        profiles: list[NutritionProfile] = []
        labels_by_product: dict[int, dict[str, Any]] = {}
        for meta in row_meta:
            label = self._label_from_row(meta["row"])
            if not self._label_has_content(label):
                continue
            product = meta["product"]
            profiles.append(
                NutritionProfile(
                    product=product,
                    serving_size=str(label.get("serving_size") or "")[:80],
                    serving_basis=str(label.get("serving_basis") or "")[:80],
                    headline=str(label.get("headline") or "")[:200],
                    note=str(label.get("note") or "")[:400],
                    guidance=str(label.get("guidance") or "")[:400],
                    nutritionist_note=str(label.get("nutritionist_note") or "")[:400],
                    hidden_sugars_found=int(label.get("hidden_sugars_found") or 0),
                    banned_ingredients_found=int(label.get("banned_ingredients_found") or 0),
                    shares_printed=bool(label.get("shares_printed", True)),
                    sugar_source=str(label.get("sugar_source") or "")[:200],
                )
            )
            labels_by_product[product.id] = label

        if not profiles:
            return

        NutritionProfile.objects.bulk_create(profiles, batch_size=self.BATCH_SIZE)
        profile_by_product = {
            profile.product_id: profile
            for profile in NutritionProfile.objects.filter(product_id__in=product_ids)
        }

        facts: list[NutritionFact] = []
        ingredients: list[Ingredient] = []
        allergens: list[Allergen] = []
        for product_id, label in labels_by_product.items():
            profile = profile_by_product.get(product_id)
            if profile is None:
                continue
            for index, fact in enumerate(label.get("facts") or []):
                name = str(fact.get("name") or "").strip()
                if not name:
                    continue
                level = str(fact.get("level") or "").strip()
                facts.append(
                    NutritionFact(
                        profile=profile,
                        name=name[:80],
                        amount=str(fact.get("amount") or "")[:40],
                        unit=str(fact.get("unit") or "")[:24],
                        daily_value=str(fact.get("daily_value") or "")[:24],
                        is_highlight=bool(fact.get("is_highlight", False)),
                        level=level,
                        note=str(fact.get("note") or "")[:200],
                        group=str(fact.get("group") or "")[:80],
                        is_subfact=bool(fact.get("is_subfact", False)),
                        sort_order=index,
                    )
                )
            for index, ingredient in enumerate(label.get("ingredients") or []):
                name = str(ingredient.get("name") or "").strip()
                if not name:
                    continue
                share = ingredient.get("share_percent")
                ingredients.append(
                    Ingredient(
                        product_id=product_id,
                        name=name[:160],
                        share_percent=Decimal(str(share)) if share not in (None, "") else None,
                        detail=str(ingredient.get("detail") or "")[:200],
                        is_flagged=bool(ingredient.get("is_flagged", False)),
                        sort_order=index,
                    )
                )
            for index, allergen in enumerate(label.get("allergens") or []):
                name = str(allergen.get("name") or "").strip()
                if not name:
                    continue
                allergens.append(
                    Allergen(
                        product_id=product_id,
                        name=name[:80],
                        detail=str(allergen.get("detail") or "")[:160],
                        sort_order=index,
                    )
                )

        if facts:
            NutritionFact.objects.bulk_create(facts, batch_size=self.BATCH_SIZE)
        if ingredients:
            Ingredient.objects.bulk_create(ingredients, batch_size=self.BATCH_SIZE)
        if allergens:
            Allergen.objects.bulk_create(allergens, batch_size=self.BATCH_SIZE)

    def _ensure_categories(self) -> dict[str, Category]:
        slugs = [slug for _name, slug, _sort in AISLE_TO_CATEGORY.values()]
        existing = {category.slug: category for category in Category.objects.filter(slug__in=slugs)}
        to_create = [
            Category(name=name, slug=slug, sort_order=sort_order, is_active=True)
            for _aisle, (name, slug, sort_order) in AISLE_TO_CATEGORY.items()
            if slug not in existing
        ]
        if to_create:
            Category.objects.bulk_create(to_create, batch_size=self.BATCH_SIZE)
            existing = {category.slug: category for category in Category.objects.filter(slug__in=slugs)}
        return existing

    def _import_lab_report(self, product: Product, row: dict[str, Any]) -> None:
        url = self._cell_text(row.get("lab_report_url"))
        if not url:
            return
        if not self._is_media_ref(url):
            return
        from catalog.image_fetch import ImageFetchError, load_report

        try:
            content, filename, _byte_size = load_report(url, base_dir=self._media_root)
        except ImageFetchError:
            return
        for old in product.lab_reports.all():
            if old.pdf:
                old.pdf.delete(save=False)
        product.lab_reports.all().delete()

        from datetime import date

        report = LabReport(
            product=product,
            lab_name=(self._cell_text(row.get("lab_report_name")) or "Lab report")[:200],
            accreditation="",
            tested_on=self._optional_date(row.get("lab_report_date")) or date.today(),
            summary=(self._cell_text(row.get("lab_report_summary")) or "Lab report attached.")[:300],
            is_current=False,
        )
        report.pdf.save(filename, content, save=False)
        report.save()
        publish_report(report)

    def _set_category_image(self, category: Category, raw_url: str) -> bool:
        content, filename, _byte_size = load_image(raw_url, base_dir=self._media_root)
        if category.image:
            category.image.delete(save=False)
        category.image.save(filename, content, save=True)
        return True

    def _import_images(self, product: Product, row: dict[str, Any]) -> None:
        urls = []
        main_image = self._cell_text(row.get("main_image"))
        if main_image and self._is_media_ref(main_image):
            urls.append(main_image)
        for raw in split_image_urls(self._cell_text(row.get("gallery"))):
            if self._is_media_ref(raw):
                urls.append(raw)
        if not urls:
            return

        product.images.all().delete()
        for index, raw_url in enumerate(urls):
            content, filename, byte_size = load_image(raw_url, base_dir=self._media_root)
            image = ProductImage(
                product=product,
                alt=product.title[:200],
                original_name=filename[:255],
                byte_size=byte_size,
                sort_order=index,
                role="primary" if index == 0 else "secondary",
            )
            image.file.save(filename, content, save=True)
        sync_image_order(product)

    def _label_from_row(self, row: dict[str, Any]) -> dict[str, Any]:
        template = self._cell_text(row.get("label_template"))
        facts = self._facts_from_row(row, template)
        return {
            "serving_size": self._serving_size(row),
            "serving_basis": self._serving_basis(row),
            "headline": self._cell_text(row.get("headline")),
            "note": self._cell_text(row.get("nutrition_note")),
            "guidance": self._cell_text(row.get("how_to_use")),
            "nutritionist_note": self._cell_text(row.get("ingredient_note")),
            "hidden_sugars_found": self._count_from_check(row.get("chk_hidden_result")),
            "banned_ingredients_found": self._count_from_check(row.get("chk_banned_result")),
            "shares_printed": self._yes_no(row.get("shares_printed"), default=True),
            "sugar_source": self._cell_text(row.get("ingredient_note")),
            "facts": facts,
            "ingredients": self._ingredients_from_row(row),
            "allergens": self._allergens_from_row(row),
        }

    def _facts_from_row(self, row: dict[str, Any], template: str) -> list[dict[str, Any]]:
        highlights = {"energy_kcal"}
        highlights.update(LABEL_TEMPLATE_HIGHLIGHTS.get(template, []))
        facts: list[dict[str, Any]] = []
        for key, (name, unit, level_key) in NUTRITION_FIELDS.items():
            amount = self._nutrition_amount(row.get(key))
            if amount is None:
                continue
            level = ""
            if level_key:
                level = LEVEL_MAP.get(self._cell_text(row.get(level_key)).lower(), "")
            facts.append(
                {
                    "name": name,
                    "amount": amount,
                    "unit": unit,
                    "is_highlight": key in highlights,
                    "level": level,
                    "is_subfact": key == "added_sugars_g",
                    "group": "Carbohydrate" if key == "added_sugars_g" else "",
                }
            )
        return facts

    def _ingredients_from_row(self, row: dict[str, Any]) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        for index in range(1, 9):
            name = self._cell_text(row.get(f"ing{index}_name"))
            if not name:
                continue
            share = row.get(f"ing{index}_pct")
            items.append(
                {
                    "name": name,
                    "share_percent": self._optional_decimal(share),
                    "detail": "",
                    "is_flagged": False,
                }
            )
        if items:
            return items

        full = self._cell_text(row.get("ingredients_full"))
        if not full:
            return []
        return [{"name": part.strip(), "share_percent": None, "detail": "", "is_flagged": False} for part in full.split(",") if part.strip()]

    def _allergens_from_row(self, row: dict[str, Any]) -> list[dict[str, Any]]:
        allergens = []
        contains = self._cell_text(row.get("allergens"))
        if contains:
            allergens.extend({"name": part.strip(), "detail": ""} for part in re.split(r"[,;/]", contains) if part.strip())
        may_contain = self._cell_text(row.get("may_contain"))
        if may_contain:
            allergens.append({"name": "May contain", "detail": may_contain})
        return allergens

    def _bulk_link_parent_groups(self, groups: dict[str, list[str]]) -> None:
        unique_slugs: list[str] = []
        for slugs in groups.values():
            seen: set[str] = set()
            ordered = []
            for slug in slugs:
                if slug and slug not in seen:
                    seen.add(slug)
                    ordered.append(slug)
            if len(ordered) < 2:
                continue
            unique_slugs.extend(ordered)

        if not unique_slugs:
            return

        products_by_slug = {product.slug: product for product in Product.objects.filter(slug__in=unique_slugs)}
        product_ids = [product.id for product in products_by_slug.values()]
        RelatedProduct.objects.filter(product_id__in=product_ids, kind=RelatedKind.FLAVOR).delete()

        links: list[RelatedProduct] = []
        for slugs in groups.values():
            unique = []
            for slug in slugs:
                if slug and slug not in unique:
                    unique.append(slug)
            if len(unique) < 2:
                continue
            for index, slug in enumerate(unique):
                product = products_by_slug.get(slug)
                if product is None:
                    continue
                for sort_order, related_slug in enumerate(other for other in unique if other != slug):
                    related = products_by_slug.get(related_slug)
                    if related is None:
                        continue
                    links.append(
                        RelatedProduct(
                            product=product,
                            related=related,
                            kind=RelatedKind.FLAVOR,
                            sort_order=sort_order,
                        )
                    )

        if links:
            RelatedProduct.objects.bulk_create(links, batch_size=self.BATCH_SIZE, ignore_conflicts=True)

    def _product_title(self, row: dict[str, Any]) -> str:
        name = self._cell_text(row.get("product_name"))
        variant = self._cell_text(row.get("variant"))
        if variant:
            return f"{name}, {variant}"[:200]
        return name[:200]

    def _product_slug(self, row: dict[str, Any]) -> str:
        sku = self._cell_text(row.get("sortd_sku"))
        if sku:
            return slugify(sku.lower())[:220]
        return slugify(self._product_title(row))[:220]

    def _variant_title(self, row: dict[str, Any]) -> str:
        pack_line = self._cell_text(row.get("pack_line"))
        if pack_line:
            return pack_line[:120]
        net_qty = self._cell_text(row.get("net_qty"))
        unit = self._cell_text(row.get("unit"))
        if net_qty and unit:
            return f"{net_qty}{unit}"[:120]
        if net_qty:
            return net_qty[:120]
        return "Default"

    def _serving_size(self, row: dict[str, Any]) -> str:
        serving = self._cell_text(row.get("serving"))
        unit = self._cell_text(row.get("unit"))
        if serving and unit and unit not in serving:
            return f"{serving}{unit}"[:80]
        return serving[:80]

    def _serving_basis(self, row: dict[str, Any]) -> str:
        basis = self._cell_text(row.get("basis"))
        printed_per = self._cell_text(row.get("printed_per"))
        if basis and printed_per:
            return f"{basis} · per {printed_per}g"[:80]
        return basis[:80]

    def _excel_status(self, row: dict[str, Any]) -> str | None:
        raw = self._cell_text(row.get("status")).lower()
        if not raw:
            return ProductStatus.DRAFT
        if raw in EXCEL_STATUS_TO_DB:
            return EXCEL_STATUS_TO_DB[raw]
        if raw in {"draft", "needs pack data", "needs price", "ready for review", "not collected"}:
            return ProductStatus.DRAFT
        return None

    def _is_product_row(self, row: dict[str, Any]) -> bool:
        sku = self._cell_text(row.get("sortd_sku"))
        if not sku:
            return False
        if sku.upper() in SKIP_IMPORT_SKUS:
            return False
        lowered = sku.lower()
        if any(marker in lowered for marker in INSTRUCTION_MARKERS):
            return False
        return bool(SKU_RE.match(sku))

    def _add_issue(self, result: ExcelImportResult, row: int, sku: str, field: str, message: str) -> None:
        result.valid = False
        result.issues.append(ImportIssue(row=row, sku=sku, field=field, message=message, level="error"))

    def _add_warning(self, result: ExcelImportResult, row: int, sku: str, field: str, message: str) -> None:
        result.warnings.append(ImportIssue(row=row, sku=sku, field=field, message=message, level="warning"))

    def _validate_image_reference(
        self,
        result: ExcelImportResult,
        row: int,
        sku: str,
        field: str,
        raw: str,
        *,
        required: bool,
        label: str | None = None,
    ) -> None:
        prefix = f"{label}: " if label else ""
        if not self._is_media_ref(raw):
            message = f"{prefix}Use a public https URL or a path relative to the workbook folder."
            if required:
                self._add_issue(result, row, sku, field, message)
            else:
                self._add_warning(result, row, sku, field, message)
            return
        if is_remote_ref(raw):
            try:
                normalize_image_url(raw)
            except ImageFetchError as exc:
                message = f"{prefix}{exc}"
                if required:
                    self._add_issue(result, row, sku, field, message)
                else:
                    self._add_warning(result, row, sku, field, message)

    def _is_media_ref(self, raw: str) -> bool:
        return is_media_ref(raw, base_dir=self._media_root)

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _has_price(value: Any) -> bool:
        if value in (None, ""):
            return False
        try:
            return Decimal(str(value)) >= 0
        except (InvalidOperation, TypeError):
            return False

    @staticmethod
    def _decimal_or_zero(value: Any) -> Decimal:
        if not value and value != 0:
            return Decimal("0.00")
        return Decimal(str(value))

    @staticmethod
    def _optional_decimal(value: Any) -> Decimal | None:
        if value in (None, ""):
            return None
        try:
            return Decimal(str(value))
        except (InvalidOperation, TypeError):
            return None

    @staticmethod
    def _optional_positive_int(value: Any) -> int | None:
        if value in (None, ""):
            return None
        try:
            parsed = int(value)
        except (TypeError, ValueError):
            return None
        return parsed if parsed > 0 else None

    @staticmethod
    def _optional_date(value: Any):
        from datetime import date, datetime

        if value in (None, ""):
            return None
        if hasattr(value, "date"):
            return value.date()
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _nutrition_amount(value: Any) -> str | None:
        if value in (None, ""):
            return None
        text = str(value).strip()
        return text[:40] if text else None

    @staticmethod
    def _yes_no(value: Any, *, default: bool) -> bool:
        text = str(value or "").strip().lower()
        if not text:
            return default
        if text in {"yes", "y", "true", "1"}:
            return True
        if text in {"no", "n", "false", "0"}:
            return False
        return default

    @staticmethod
    def _count_from_check(value: Any) -> int:
        text = str(value or "").strip()
        if not text:
            return 0
        if text.isdigit():
            return int(text)
        return 0

    @staticmethod
    def _label_has_content(label: dict[str, Any]) -> bool:
        if any(
            label.get(key)
            for key in (
                "serving_size",
                "serving_basis",
                "headline",
                "note",
                "guidance",
                "nutritionist_note",
                "sugar_source",
            )
        ):
            return True
        return bool(label.get("facts") or label.get("ingredients") or label.get("allergens"))
