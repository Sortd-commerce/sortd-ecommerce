from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from catalog.excel_importer import ExcelCatalogImporter
from catalog.models import Category, Product, ProductVariant
from django.test import TestCase
from openpyxl import Workbook


def build_workbook(rows: list[dict[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Launch range"
    keys = [
        "sortd_sku",
        "status",
        "product_name",
        "brand",
        "aisle",
        "shelf",
        "tags",
        "price",
        "pack_line",
        "max_order",
        "short_desc",
        "main_image",
        "gallery",
        "headline",
        "basis",
        "serving",
        "printed_per",
        "protein_g",
        "label_template",
        "ing1_name",
        "ing1_pct",
        "allergens",
        "shares_printed",
    ]
    sheet.append(["Section"] + [""] * (len(keys) - 1))
    sheet.append(keys)
    sheet.append(keys)
    sheet.append(["instructions"] * len(keys))
    for row in rows:
        sheet.append([row.get(key, "") for key in keys])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ExcelImporterTests(TestCase):
    def test_validate_flags_duplicate_skus(self):
        payload = build_workbook(
            [
                {"sortd_sku": "SRT-BRK-001", "product_name": "Spread", "brand": "Sortd", "aisle": "Breakfast & spreads"},
                {"sortd_sku": "SRT-BRK-001", "product_name": "Other", "brand": "Sortd", "aisle": "Breakfast & spreads"},
            ]
        )
        importer = ExcelCatalogImporter()
        rows = importer.parse_workbook(BytesIO(payload))
        result = importer.validate_rows(rows)
        self.assertFalse(result.valid)
        self.assertTrue(any("Duplicate" in issue.message for issue in result.errors))

    def test_import_creates_draft_product_without_price(self):
        payload = build_workbook(
            [
                {
                    "sortd_sku": "SRT-BRK-010",
                    "status": "Needs pack data",
                    "product_name": "Almond Spread",
                    "brand": "Sortd",
                    "aisle": "Breakfast & spreads",
                    "pack_line": "200g jar",
                    "ing1_name": "Almonds",
                    "label_template": "Nut butters & spreads",
                    "protein_g": "6",
                }
            ]
        )
        importer = ExcelCatalogImporter()
        with patch("catalog.excel_importer.fetch_image") as fetch_image:
            result = importer.import_file(BytesIO(payload), dry_run=False)
        self.assertTrue(result.valid)
        self.assertEqual(result.created, 1)
        product = Product.objects.get(slug="srt-brk-010")
        self.assertEqual(product.status, "draft")
        variant = ProductVariant.objects.get(sku="SRT-BRK-010")
        self.assertEqual(variant.title, "200g jar")
        self.assertEqual(variant.price, Decimal("0.00"))
        fetch_image.assert_not_called()

    def test_import_downloads_main_and_gallery_images(self):
        payload = build_workbook(
            [
                {
                    "sortd_sku": "SRT-SNK-010",
                    "status": "Needs pack data",
                    "product_name": "Protein Bar",
                    "brand": "Sortd",
                    "aisle": "Snacks & bars",
                    "pack_line": "Single bar",
                    "main_image": "https://example.com/main.jpg",
                    "gallery": "https://example.com/one.jpg, https://example.com/two.jpg",
                    "ing1_name": "Oats",
                    "label_template": "Protein & snack bars",
                }
            ]
        )
        importer = ExcelCatalogImporter()
        with patch.object(importer, "_import_images") as import_images:
            result = importer.import_file(BytesIO(payload), dry_run=False)
        self.assertTrue(result.valid)
        product = Product.objects.get(slug="srt-snk-010")
        import_images.assert_called_once()
        self.assertEqual(import_images.call_args[0][0].pk, product.pk)
        self.assertTrue(Category.objects.filter(slug="snacks-bars").exists())

    def test_import_updates_aisle_images_from_aisles_sheet(self):
        workbook = Workbook()
        products = workbook.active
        products.title = "Launch range"
        keys = ["sortd_sku", "product_name", "brand", "aisle", "label_template", "ing1_name"]
        products.append(["Section"] + [""] * (len(keys) - 1))
        products.append(keys)
        products.append(keys)
        products.append(["instructions"] * len(keys))
        products.append(["SRT-BRK-011", "Oats", "Sortd", "Breakfast & spreads", "Breakfast & oats", "Oats"])

        aisles = workbook.create_sheet("Aisles")
        aisles.append(["aisle", "image_url"])
        aisles.append(["Breakfast & spreads", "https://example.com/breakfast.jpg"])

        buffer = BytesIO()
        workbook.save(buffer)
        importer = ExcelCatalogImporter()
        with patch.object(importer, "_set_category_image", return_value=True) as set_image:
            result = importer.import_file(BytesIO(buffer.getvalue()), dry_run=False)
        self.assertTrue(result.valid)
        self.assertEqual(result.aisle_images_updated, 1)
        set_image.assert_called_once()
        category = Category.objects.get(slug="breakfast-spreads")
        self.assertEqual(set_image.call_args[0][0].pk, category.pk)
        self.assertEqual(set_image.call_args[0][1], "https://example.com/breakfast.jpg")

    def test_import_stores_brand_shelf_tags_and_max_order(self):
        payload = build_workbook(
            [
                {
                    "sortd_sku": "SRT-CHO-020",
                    "status": "Needs pack data",
                    "product_name": "Dark Bar",
                    "brand": "PLAAAY",
                    "aisle": "Chocolate",
                    "shelf": "Dark chocolate",
                    "tags": "vegan, dark",
                    "max_order": "4",
                    "pack_line": "40 g",
                    "ing1_name": "Cacao",
                    "label_template": "Chocolate",
                }
            ]
        )
        importer = ExcelCatalogImporter()
        with patch.object(importer, "_import_images"):
            result = importer.import_file(BytesIO(payload), dry_run=False)
        self.assertTrue(result.valid)
        product = Product.objects.get(slug="srt-cho-020")
        self.assertEqual(product.brand, "PLAAAY")
        self.assertEqual(product.shelf, "Dark chocolate")
        self.assertEqual(product.tags, "vegan, dark")
        variant = ProductVariant.objects.get(sku="SRT-CHO-020")
        self.assertEqual(variant.max_order, 4)
