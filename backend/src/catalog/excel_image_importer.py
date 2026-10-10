"""Upload a small spreadsheet batch of external product images."""

from __future__ import annotations

import csv
import ipaddress
import socket
from io import BytesIO, StringIO
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urljoin, urlparse
from urllib.request import HTTPRedirectHandler, Request, build_opener

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import UploadedFile
from django.db import transaction
from django.db.models import Max
from openpyxl import load_workbook

from catalog.images import sync_image_order
from catalog.models import ImageRole, Product, ProductImage
from catalog.schemas import serialize_image
from core.uploads import MAX_IMAGE_BYTES

MAX_IMAGE_BATCH_ROWS = 10
MAX_REDIRECTS = 4
IMAGE_CONTENT_TYPES = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
}


class ProductImageSheetError(ValueError):
    pass


class _NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def _public_https_url(raw_url: str) -> str:
    url = str(raw_url or "").strip()
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Image URL must be a public HTTPS URL.")
    host = parsed.hostname.rstrip(".").lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith((".local", ".internal", ".localhost")):
        raise ValueError("Local and internal image hosts are not allowed.")
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(host, parsed.port or 443, type=socket.SOCK_STREAM)}
    except OSError as exc:
        raise ValueError("Image host could not be resolved.") from exc
    if not addresses:
        raise ValueError("Image host could not be resolved.")
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%", 1)[0])
        if not ip.is_global:
            raise ValueError("Image URL must resolve to a public internet address.")
    return url


def _download_public_image(raw_url: str) -> tuple[ContentFile, str, int]:
    url = _public_https_url(raw_url)
    opener = build_opener(_NoRedirect)
    for redirect_count in range(MAX_REDIRECTS + 1):
        request = Request(url, headers={"User-Agent": "SortdCatalogImageUploader/1.0"})
        try:
            response = opener.open(request, timeout=20)
        except HTTPError as exc:
            if exc.code in {301, 302, 303, 307, 308}:
                location = exc.headers.get("Location")
                if not location or redirect_count >= MAX_REDIRECTS:
                    exc.close()
                    raise ValueError("Image URL redirected too many times.") from exc
                url = _public_https_url(urljoin(url, location))
                exc.close()
                continue
            exc.close()
            raise ValueError(f"Image server returned HTTP {exc.code}.") from exc
        except Exception as exc:
            raise ValueError("Could not download image from the supplied URL.") from exc

        with response:
            payload = response.read(MAX_IMAGE_BYTES + 1)
            if len(payload) > MAX_IMAGE_BYTES:
                raise ValueError("Image exceeds the 5 MB size limit.")
            content_type = (response.headers.get("Content-Type") or "").split(";", 1)[0].strip().lower()
        extension = IMAGE_CONTENT_TYPES.get(content_type)
        if not extension:
            raise ValueError("URL did not return a JPEG, PNG, or WebP image.")
        stem = Path(urlparse(url).path).stem or "product-image"
        safe_stem = "".join(char if char.isalnum() or char in "-_" else "-" for char in stem)[:70].strip("-") or "product-image"
        filename = f"{safe_stem}{extension}"
        return ContentFile(payload, name=filename), filename, len(payload)
    raise ValueError("Image URL redirected too many times.")


def _read_rows(file_obj: UploadedFile) -> list[dict]:
    name = (getattr(file_obj, "name", "") or "").lower()
    content = file_obj.read()
    file_obj.seek(0)
    if name.endswith(".xlsx"):
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower() for value in next(iterator, ())]
        records = [dict(zip(headers, values)) for values in iterator if any(value not in (None, "") for value in values)]
        workbook.close()
        return records
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ProductImageSheetError("CSV must use UTF-8 encoding.") from exc
        return [dict(row) for row in csv.DictReader(StringIO(text)) if any(str(value or "").strip() for value in row.values())]
    raise ProductImageSheetError("Upload a .xlsx or .csv file.")


def process_product_image_sheet(file_obj: UploadedFile) -> dict:
    """Process at most ten rows. Each non-empty input row describes one image."""
    raw_rows = _read_rows(file_obj)
    if not raw_rows:
        raise ProductImageSheetError("The sheet has no image rows.")
    if len(raw_rows) > MAX_IMAGE_BATCH_ROWS:
        raise ProductImageSheetError(f"Upload no more than {MAX_IMAGE_BATCH_ROWS} image rows at a time.")

    if not raw_rows[0]:
        raise ProductImageSheetError("The sheet must have a header row.")
    header_names = {str(key).strip().lower() for key in raw_rows[0] if key}
    if "product_id" not in header_names or not ({"image_url", "external_image_url"} & header_names):
        raise ProductImageSheetError("Required columns: product_id and image_url (or external_image_url).")

    results = []
    for row_number, row in enumerate(raw_rows, start=2):
        normalized = {str(key).strip().lower(): value for key, value in row.items() if key}
        product_id_value = str(normalized.get("product_id") or "").strip()
        image_url = str(normalized.get("image_url") or normalized.get("external_image_url") or "").strip()
        alt = str(normalized.get("alt") or "").strip()
        outcome = {
            "row": row_number,
            "product_id": product_id_value,
            "image_url": image_url,
            "cloudinary_url": "",
            "status": "error",
            "error": "",
        }
        try:
            product_id = int(product_id_value)
            if product_id < 1:
                raise ValueError("Product ID must be a positive integer.")
            if not image_url:
                raise ValueError("Image URL is required.")
            product = Product.objects.filter(pk=product_id).first()
            if product is None:
                raise ValueError(f"Product {product_id} was not found.")

            content, filename, byte_size = _download_public_image(image_url)
            next_order = (product.images.aggregate(max_order=Max("sort_order"))["max_order"] or -1) + 1
            image = ProductImage(
                product=product,
                alt=(alt or product.title)[:200],
                original_name=filename[:255],
                byte_size=byte_size,
                sort_order=next_order,
                role=ImageRole.PRIMARY if next_order == 0 else ImageRole.SECONDARY,
            )
            image.file = content
            image.full_clean()
            with transaction.atomic():
                image.save()
                sync_image_order(product)
            image.refresh_from_db()
            outcome.update(
                cloudinary_url=serialize_image(image)["url"],
                status="uploaded",
            )
        except (TypeError, ValueError) as exc:
            outcome["error"] = str(exc)
        except DjangoValidationError as exc:
            outcome["error"] = "; ".join(exc.messages)
        except Exception as exc:
            outcome["error"] = f"Upload failed: {exc}"
        results.append(outcome)

    return {"row_count": len(results), "uploaded": sum(row["status"] == "uploaded" for row in results), "results": results}
