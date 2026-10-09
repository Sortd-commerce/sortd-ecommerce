"""Build a client-facing import workbook from the full internal listing sheet."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from catalog.import_columns import IMPORT_COLUMNS, LIST_VALUES

SHEET_NAME = "Launch range"
SECTION_FILL = PatternFill("solid", fgColor="D9D2C6")
LABEL_FILL = PatternFill("solid", fgColor="F3F0EA")
KEY_FILL = PatternFill("solid", fgColor="EDE8E0")
HINT_FILL = PatternFill("solid", fgColor="FAF8F5")
REQUIRED_FILL = PatternFill("solid", fgColor="E6ECE6")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0C8BC"),
    right=Side(style="thin", color="D0C8BC"),
    top=Side(style="thin", color="D0C8BC"),
    bottom=Side(style="thin", color="D0C8BC"),
)


def build_import_workbook(source_path: Path, output_path: Path) -> int:
    """Write a simplified workbook and return the number of product rows copied."""
    from catalog.excel_importer import ExcelCatalogImporter

    with source_path.open("rb") as handle:
        rows = ExcelCatalogImporter().parse_workbook(handle)
    return build_import_workbook_from_rows(rows, output_path)


def build_import_workbook_from_rows(
    rows: list[dict[str, Any]],
    output_path: Path,
    *,
    preserve_image_paths: bool = False,
    aisle_rows: list[dict[str, str]] | None = None,
) -> int:
    """Write a simplified import workbook from parsed row dicts."""
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Read me"
    _write_readme(readme)

    listing = workbook.create_sheet(SHEET_NAME)
    row_count = _write_listing_sheet(listing, rows, preserve_image_paths=preserve_image_paths)

    lists = workbook.create_sheet("Lists")
    _write_lists_sheet(lists)

    aisles = workbook.create_sheet("Aisles")
    _write_aisles_sheet(aisles, aisle_rows=aisle_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return row_count


def _write_readme(sheet) -> None:
    lines = [
        ("Sortd product import sheet", ""),
        ("", ""),
        ("Who fills this", "Sortd team or supplier contact. One row = one product SKU."),
        ("Sheet to fill", "Launch range - product rows start on row 5."),
        ("Header rows", "Row 1 = section, row 2 = column names, row 3 = system keys (do not edit), row 4 = guidance."),
        ("Before Live", "Set status to Live only when price, copy, pack line, image URLs, nutrition, and ingredients are complete."),
        ("Images", "Use public https links in Main image URL, Gallery, and the Aisles sheet. Google Drive share links work."),
        ("Aisles sheet", "One row per shop aisle with a hero image URL for the storefront picker."),
        ("Import", "Sortd admin > Catalog > Import Excel."),
        ("", ""),
        ("Required always", "Sortd SKU, Status, Product name, Brand, Aisle, Label template."),
        ("Required for Live", "Pack line, Price, Short description, Label headline, Main image URL, Label basis, Serving size, Panel printed per, Ingredients, Shares printed, Allergens."),
    ]
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 92
    for index, (left, right) in enumerate(lines, start=1):
        left_cell = sheet.cell(row=index, column=1, value=left)
        sheet.cell(row=index, column=2, value=right)
        if index == 1:
            left_cell.font = Font(bold=True, size=14)
        elif left:
            left_cell.font = Font(bold=True)


def _write_listing_sheet(sheet, rows: list[dict[str, Any]], *, preserve_image_paths: bool = False) -> int:
    column_count = len(IMPORT_COLUMNS)
    section_row = [column.section for column in IMPORT_COLUMNS]
    label_row = [column.label for column in IMPORT_COLUMNS]
    key_row = [column.key for column in IMPORT_COLUMNS]
    hint_row = [column.hint or "Optional." for column in IMPORT_COLUMNS]

    for col_index, value in enumerate(section_row, start=1):
        sheet.cell(row=1, column=col_index, value=value)
    for col_index, value in enumerate(label_row, start=1):
        sheet.cell(row=2, column=col_index, value=value)
    for col_index, value in enumerate(key_row, start=1):
        sheet.cell(row=3, column=col_index, value=value)
    for col_index, value in enumerate(hint_row, start=1):
        sheet.cell(row=4, column=col_index, value=value)

    for source in rows:
        values = []
        for column in IMPORT_COLUMNS:
            value = source.get(column.key)
            if (
                not preserve_image_paths
                and column.key in {"main_image", "gallery"}
                and value
                and not str(value).strip().startswith("http")
            ):
                value = ""
            values.append(value)
        sheet.append(values)

    _merge_section_headers(sheet, column_count)
    _style_header_rows(sheet, column_count)
    _apply_validations(sheet, len(rows))
    _autosize_columns(sheet, column_count)
    sheet.freeze_panes = "A5"
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 36
    sheet.row_dimensions[3].height = 18
    sheet.row_dimensions[4].height = 48
    return len(rows)


def _merge_section_headers(sheet, column_count: int) -> None:
    spans: list[tuple[str, int, int]] = []
    current = IMPORT_COLUMNS[0].section
    start = 1
    for index, column in enumerate(IMPORT_COLUMNS[1:], start=2):
        if column.section != current:
            spans.append((current, start, index - 1))
            current = column.section
            start = index
    spans.append((current, start, column_count))

    for title, start_col, end_col in spans:
        sheet.cell(row=1, column=start_col, value=title)
        if end_col > start_col:
            sheet.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = sheet.cell(row=1, column=start_col)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_header_rows(sheet, column_count: int) -> None:
    for col_index, column in enumerate(IMPORT_COLUMNS, start=1):
        section_cell = sheet.cell(row=1, column=col_index)
        section_cell.font = Font(bold=True, size=11)
        section_cell.fill = SECTION_FILL
        section_cell.border = THIN_BORDER

        label_cell = sheet.cell(row=2, column=col_index)
        label_cell.font = Font(bold=True, size=10)
        label_cell.fill = REQUIRED_FILL if column.required in {"always", "live"} else LABEL_FILL
        label_cell.alignment = Alignment(wrap_text=True, vertical="top")
        label_cell.border = THIN_BORDER

        key_cell = sheet.cell(row=3, column=col_index)
        key_cell.font = Font(size=9, italic=True, color="666666")
        key_cell.fill = KEY_FILL
        key_cell.alignment = Alignment(wrap_text=True, vertical="top")
        key_cell.border = THIN_BORDER

        hint_cell = sheet.cell(row=4, column=col_index)
        hint_cell.font = Font(size=9, color="444444")
        hint_cell.fill = HINT_FILL
        hint_cell.alignment = Alignment(wrap_text=True, vertical="top")
        hint_cell.border = THIN_BORDER


def _write_aisles_sheet(sheet, *, aisle_rows: list[dict[str, str]] | None = None) -> None:
    sheet.append(["aisle", "image_url"])
    url_by_aisle = {
        str(row.get("aisle") or "").strip(): str(row.get("image_url") or "").strip()
        for row in (aisle_rows or [])
        if str(row.get("aisle") or "").strip()
    }
    for aisle in LIST_VALUES["Aisle"]:
        sheet.append([aisle, url_by_aisle.get(aisle, "")])
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 48
    for col in (1, 2):
        header = sheet.cell(row=1, column=col)
        header.font = Font(bold=True)
        header.fill = LABEL_FILL
        header.border = THIN_BORDER
    sheet.freeze_panes = "A2"


def _write_lists_sheet(sheet) -> None:
    headers = list(LIST_VALUES.keys())
    sheet.append(headers)
    max_len = max(len(values) for values in LIST_VALUES.values())
    for index in range(max_len):
        sheet.append([values[index] if index < len(values) else "" for values in LIST_VALUES.values()])

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        sheet.column_dimensions[letter].width = 30
        header = sheet.cell(row=1, column=col)
        header.font = Font(bold=True)
        header.fill = LABEL_FILL
        header.border = THIN_BORDER
    sheet.freeze_panes = "A2"


def _apply_validations(sheet, data_rows: int) -> None:
    if data_rows < 1:
        return
    last_row = data_rows + 4
    key_index = {column.key: index + 1 for index, column in enumerate(IMPORT_COLUMNS)}
    list_columns = {name: get_column_letter(index + 1) for index, name in enumerate(LIST_VALUES.keys())}

    def add_list_validation(key: str, list_header: str) -> None:
        col = get_column_letter(key_index[key])
        list_col = list_columns[list_header]
        last_list_row = len(LIST_VALUES[list_header]) + 1
        validation = DataValidation(
            type="list",
            formula1=f"=Lists!${list_col}$2:${list_col}${last_list_row}",
            allow_blank=True,
            showDropDown=True,
        )
        validation.add(f"{col}5:{col}{last_row}")
        sheet.add_data_validation(validation)

    add_list_validation("status", "Status")
    add_list_validation("aisle", "Aisle")
    add_list_validation("label_template", "Template")
    add_list_validation("unit", "Unit")
    add_list_validation("variation_theme", "Variation")
    add_list_validation("shares_printed", "Yesno")


def _autosize_columns(sheet, column_count: int) -> None:
    for index in range(1, column_count + 1):
        letter = get_column_letter(index)
        column = IMPORT_COLUMNS[index - 1]
        label_len = len(column.label)
        hint_len = min(len(column.hint or ""), 36)
        width = min(max(label_len, hint_len, len(column.key)) + 4, 36)
        if column.key in {"short_desc", "ingredients_full", "gallery", "main_image"}:
            width = 36
        sheet.column_dimensions[letter].width = max(width, 14)
